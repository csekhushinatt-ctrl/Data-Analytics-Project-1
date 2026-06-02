from __future__ import annotations

import csv
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
RAW_PATH = ROOT / "data" / "raw_orders.csv"
OUTPUT_DIR = ROOT / "outputs"
CLEANED_PATH = OUTPUT_DIR / "cleaned_orders.csv"
AUDIT_PATH = OUTPUT_DIR / "cleaning_audit.csv"
LOG_PATH = OUTPUT_DIR / "cleaning_log.csv"
WORKBOOK_PATH = OUTPUT_DIR / "data_cleaning_project_1.xlsx"


def normalize_text(value: object, default: str | None = None) -> str | None:
    if pd.isna(value) or str(value).strip() == "":
        return default
    return re.sub(r"\s+", " ", str(value).strip()).title()


def parse_quantity(value: object) -> float | None:
    if pd.isna(value) or str(value).strip() == "":
        return None
    text = str(value).strip().lower()
    word_numbers = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5}
    if text in word_numbers:
        return float(word_numbers[text])
    parsed = pd.to_numeric(text, errors="coerce")
    if pd.isna(parsed):
        return None
    return abs(float(parsed)) if float(parsed) != 0 else None


def parse_price(value: object) -> float | None:
    if pd.isna(value) or str(value).strip() == "":
        return None
    cleaned = re.sub(r"[^0-9.\-]", "", str(value))
    parsed = pd.to_numeric(cleaned, errors="coerce")
    if pd.isna(parsed):
        return None
    return abs(float(parsed))


def parse_date_series(series: pd.Series) -> pd.Series:
    formats = [
        "%Y/%m/%d",
        "%d-%m-%Y",
        "%b %d %Y",
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%Y.%m.%d",
        "%d %b %Y",
    ]
    parsed = pd.Series(pd.NaT, index=series.index, dtype="datetime64[ns]")
    for date_format in formats:
        missing = parsed.isna()
        parsed.loc[missing] = pd.to_datetime(series.loc[missing], errors="coerce", format=date_format)
    return parsed


def invalid_date_count(series: pd.Series) -> int:
    parsed = pd.to_datetime(series, errors="coerce", format="%Y-%m-%d")
    return int(parsed.isna().sum())


def build_audit(raw: pd.DataFrame, cleaned: pd.DataFrame) -> pd.DataFrame:
    checks = [
        {
            "Check": "Total rows",
            "Before": len(raw),
            "After": len(cleaned),
            "Status": "Pass",
        },
        {
            "Check": "Duplicate Order IDs",
            "Before": int(raw["Order ID"].duplicated().sum()),
            "After": int(cleaned["Order ID"].duplicated().sum()),
            "Status": "Pass" if cleaned["Order ID"].duplicated().sum() == 0 else "Fail",
        },
        {
            "Check": "Incorrectly formatted dates",
            "Before": int(parse_date_series(raw["Order Date"]).isna().sum()),
            "After": invalid_date_count(cleaned["Order Date"]),
            "Status": "Pass" if invalid_date_count(cleaned["Order Date"]) == 0 else "Fail",
        },
        {
            "Check": "Missing critical values",
            "Before": int(raw[["Order ID", "Order Date", "Customer Name", "Quantity", "Unit Price"]].isna().sum().sum()),
            "After": int(cleaned[["Order ID", "Order Date", "Customer Name", "Quantity", "Unit Price"]].isna().sum().sum()),
            "Status": "Pass",
        },
        {
            "Check": "Non-positive quantities",
            "Before": "Mixed text/negative/zero values present",
            "After": int((cleaned["Quantity"] <= 0).sum()),
            "Status": "Pass" if (cleaned["Quantity"] <= 0).sum() == 0 else "Fail",
        },
    ]
    return pd.DataFrame(checks)


def clean_data(raw: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    working = raw.copy()
    log_rows: list[dict[str, str]] = []

    duplicate_mask = working["Order ID"].duplicated(keep="first")
    for order_id in working.loc[duplicate_mask, "Order ID"]:
        log_rows.append({"Order ID": order_id, "Field": "Order ID", "Action": "Duplicate row removed"})
    working = working.loc[~duplicate_mask].copy()

    parsed_dates = parse_date_series(working["Order Date"])
    median_date = parsed_dates.dropna().median()
    date_missing_mask = parsed_dates.isna()
    for order_id in working.loc[date_missing_mask, "Order ID"]:
        log_rows.append({"Order ID": order_id, "Field": "Order Date", "Action": f"Invalid/missing date imputed as {median_date:%Y-%m-%d}"})
    working["Order Date"] = parsed_dates.fillna(median_date).dt.strftime("%Y-%m-%d")

    for col, default in [("Customer Name", None), ("City", "Unknown"), ("Category", "Uncategorized"), ("Status", "Unknown")]:
        before = working[col].copy()
        working[col] = working[col].apply(lambda value: normalize_text(value, default))
        changed = before.astype(str).str.strip() != working[col].astype(str).str.strip()
        for order_id in working.loc[changed, "Order ID"]:
            log_rows.append({"Order ID": order_id, "Field": col, "Action": "Text standardized or missing value filled"})

    quantities = working["Quantity"].apply(parse_quantity)
    median_quantity = int(quantities.dropna().median())
    quantity_changed = quantities.isna() | (working["Quantity"].astype(str).str.strip() != quantities.fillna(median_quantity).astype(int).astype(str))
    working["Quantity"] = quantities.fillna(median_quantity).round().astype(int)
    for order_id in working.loc[quantity_changed, "Order ID"]:
        log_rows.append({"Order ID": order_id, "Field": "Quantity", "Action": f"Converted to positive integer or imputed as {median_quantity}"})

    prices = working["Unit Price"].apply(parse_price)
    median_price = round(float(prices.dropna().median()), 2)
    price_changed = prices.isna() | (working["Unit Price"].astype(str).str.strip() != prices.fillna(median_price).round(2).astype(str))
    working["Unit Price"] = prices.fillna(median_price).round(2)
    for order_id in working.loc[price_changed, "Order ID"]:
        log_rows.append({"Order ID": order_id, "Field": "Unit Price", "Action": f"Converted to numeric currency or imputed as {median_price}"})

    working["Total Amount"] = (working["Quantity"] * working["Unit Price"]).round(2)
    working = working.sort_values("Order ID").reset_index(drop=True)

    log = pd.DataFrame(log_rows).drop_duplicates()
    return working, log


def write_sheet(ws, dataframe: pd.DataFrame, title: str) -> None:
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(dataframe.columns)))
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws.append(list(dataframe.columns))
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.alignment = Alignment(horizontal="center")
    for row in dataframe.itertuples(index=False):
        ws.append(list(row))
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 45)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def export_workbook(raw: pd.DataFrame, cleaned: pd.DataFrame, audit: pd.DataFrame, log: pd.DataFrame) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, frame, title in [
        ("Audit Summary", audit, "Data Quality Audit"),
        ("Raw Data", raw, "Raw Orders Dataset"),
        ("Cleaned Data", cleaned, "Cleaned Orders Dataset"),
        ("Cleaning Log", log, "Row-Level Cleaning Log"),
    ]:
        write_sheet(wb.create_sheet(sheet_name), frame, title)
    wb.save(WORKBOOK_PATH)


def main() -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)
    raw = pd.read_csv(RAW_PATH, dtype=str, keep_default_na=False, quoting=csv.QUOTE_MINIMAL)
    raw = raw.replace("", pd.NA)
    cleaned, log = clean_data(raw)
    audit = build_audit(raw, cleaned)

    cleaned.to_csv(CLEANED_PATH, index=False)
    audit.to_csv(AUDIT_PATH, index=False)
    log.to_csv(LOG_PATH, index=False)
    export_workbook(raw, cleaned, audit, log)

    duplicate_ids = int(cleaned["Order ID"].duplicated().sum())
    bad_dates = invalid_date_count(cleaned["Order Date"])
    missing_critical = int(cleaned[["Order ID", "Order Date", "Customer Name", "Quantity", "Unit Price"]].isna().sum().sum())

    print("Cleaning complete")
    print(f"Rows before: {len(raw)}")
    print(f"Rows after: {len(cleaned)}")
    print(f"Duplicate Order IDs after cleaning: {duplicate_ids}")
    print(f"Incorrectly formatted dates after cleaning: {bad_dates}")
    print(f"Missing critical values after cleaning: {missing_critical}")
    print(f"Workbook: {WORKBOOK_PATH}")

    if duplicate_ids or bad_dates or missing_critical:
        raise SystemExit("Validation failed")


if __name__ == "__main__":
    main()
